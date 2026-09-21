import openpyxl
import datetime
import pandas as pd

file_path = 'Gold_ETF_Daily_Price_Tracker.xlsx'
wb = openpyxl.load_workbook(file_path)

ws_dash = wb["Dashboard"]
last_dash_date = None
last_dash_row = ws_dash.max_row
while last_dash_row > 1:
    val = ws_dash.cell(row=last_dash_row, column=1).value
    if val:
        if isinstance(val, datetime.datetime):
            last_dash_date = val.date()
        else:
            try:
                last_dash_date = pd.to_datetime(val).date()
            except:
                pass
    if last_dash_date:
        break
    last_dash_row -= 1

print("Dashboard last date:", last_dash_date)

ws_ref = wb['GOLDBETA']
new_dates = []
for r in range(10, ws_ref.max_row + 1):
    d_val = ws_ref.cell(row=r, column=1).value
    if d_val:
        d_parsed = None
        if isinstance(d_val, datetime.datetime):
            d_parsed = d_val.date()
        else:
            try:
                d_parsed = pd.to_datetime(d_val).date()
            except:
                pass
        if d_parsed and d_parsed > last_dash_date:
            price = ws_ref.cell(row=r, column=3).value or 0
            if float(price) > 0:
                new_dates.append(d_parsed)

print("new_dates:", new_dates)
