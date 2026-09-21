import openpyxl
file_path = 'Gold_ETF_Daily_Price_Tracker.xlsx'
wb = openpyxl.load_workbook(file_path)
ws = wb['GOLDBETA']
print("GOLDBETA rows from 18-Sep onwards:")
for r in range(max(10, ws.max_row - 10), ws.max_row + 1):
    d = ws.cell(row=r, column=1).value
    p = ws.cell(row=r, column=3).value
    print(f"Row {r}: Date={d}, Price={p}")
