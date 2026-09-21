import openpyxl
import datetime
import pandas as pd

file_path = 'Gold_ETF_Daily_Price_Tracker.xlsx'
wb = openpyxl.load_workbook(file_path)
ws = wb['GOLDBETA']

last_row = 9
last_date = None
for r in range(10, ws.max_row + 1):
    cell_val = ws.cell(row=r, column=1).value
    if isinstance(cell_val, datetime.datetime):
        last_row = r
        last_date = cell_val.date()
    elif isinstance(cell_val, str) and cell_val.strip() != "":
        try:
            last_date = pd.to_datetime(cell_val).date()
            last_row = r
        except:
            pass

print("Last date:", last_date)
print("Today:", datetime.datetime.now().date())
print("last_date >= today:", last_date >= datetime.datetime.now().date())
