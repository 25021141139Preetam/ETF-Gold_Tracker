import openpyxl
import yfinance as yf
import pandas as pd
import datetime
import logging
import requests

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

def get_yfinance_session():
    session = requests.Session()
    session.headers.update({
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/115.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8"
    })
    return session

def fetch_data_fallback(symbol, start_date, end_date):
    try:
        session = get_yfinance_session()
        ticker = yf.Ticker(symbol, session=session)
        df = ticker.history(start=start_date.strftime("%Y-%m-%d"), end=end_date.strftime("%Y-%m-%d"))
        return df
    except Exception as e:
        logging.error(f"Error fetching data for {symbol}: {e}")
        return pd.DataFrame()

def get_historical_gold(start_date, end_date):
    session = get_yfinance_session()
    gc = yf.Ticker("GC=F", session=session)
    inr = yf.Ticker("INR=X", session=session)
    
    s_str = (start_date - datetime.timedelta(days=5)).strftime("%Y-%m-%d")
    e_str = (end_date + datetime.timedelta(days=2)).strftime("%Y-%m-%d")
    
    df_gc = gc.history(start=s_str, end=e_str)
    df_inr = inr.history(start=s_str, end=e_str)
    
    df_gc.index = pd.to_datetime(df_gc.index).tz_localize(None).normalize()
    df_inr.index = pd.to_datetime(df_inr.index).tz_localize(None).normalize()
    
    df = pd.merge(df_gc[['Close']], df_inr[['Close']], left_index=True, right_index=True, how='outer', suffixes=('_GC', '_INR'))
    df = df.ffill().bfill()
    df['Domestic_10g'] = (df['Close_GC'] * df['Close_INR']) / 3.11034768 * 1.15 * 1.03
    return df

def update_tracker(file_path):
    logging.info(f"Opening workbook: {file_path}")
    try:
        wb = openpyxl.load_workbook(file_path)
    except Exception as e:
        logging.error(f"Failed to open workbook: {e}")
        return

    etf_sheets = ['GOLDBETA', 'HDFCGOLD', 'GOLDIETF', 'BSLGOLDETF', 'LICMFGOLD']
    today = datetime.datetime.now().date()
    end_date_yf = today + datetime.timedelta(days=1)

    for sheet_name in etf_sheets:
        if sheet_name not in wb.sheetnames:
            continue
            
        ws = wb[sheet_name]
        
        last_row = 9
        last_date = None
        for r in range(10, ws.max_row + 1):
            cell_val = ws.cell(row=r, column=1).value
            if isinstance(cell_val, datetime.datetime):
                last_row = r
                last_date = cell_val.date()
            elif isinstance(cell_val, str) and cell_val.strip() != "":
                try:
                    last_date = pd.to_datetime(cell_val).date()
                    last_row = r
                except:
                    pass
        
        if last_date is None:
            continue
            
        if last_date >= today:
            continue
            
        start_date_fetch = last_date + datetime.timedelta(days=1)
        yf_symbol = f"{sheet_name}.NS"
        
        logging.info(f"[{sheet_name}] Fetching {yf_symbol} from {start_date_fetch} to {today}")
        df = fetch_data_fallback(yf_symbol, start_date_fetch, end_date_yf)
        
        if df.empty:
            continue
            
        curr_row = last_row + 1
        for date_idx, row_data in df.iterrows():
            d = date_idx.date()
            if d <= last_date:
                continue
                
            close_price = round(row_data['Close'], 2)
            day_str = d.strftime("%a")
            
            ws.cell(row=curr_row, column=1).value = pd.to_datetime(d) 
            ws.cell(row=curr_row, column=1).number_format = 'yyyy-mm-dd'
            ws.cell(row=curr_row, column=2).value = day_str 
            ws.cell(row=curr_row, column=3).value = close_price 
            ws.cell(row=curr_row, column=10).value = "Available" 
            curr_row += 1

    if "Dashboard" in wb.sheetnames:
        ws_dash = wb["Dashboard"]
        etfs_dash = {
            'GOLDBETA': (2, 3, 2356),
            'HDFCGOLD': (4, 5, 1934),
            'GOLDIETF': (6, 7, 1544),
            'BSLGOLDETF': (8, 9, 1111),
            'LICMFGOLD': (10, 11, 733)
        }
        
        last_dash_date = None
        last_dash_row = ws_dash.max_row
        while last_dash_row > 1:
            val = ws_dash.cell(row=last_dash_row, column=1).value
            if val:
                if isinstance(val, datetime.datetime):
                    last_dash_date = val.date()
                else:
                    try:
                        last_dash_date = pd.to_datetime(val).date()
                    except:
                        pass
            if last_dash_date:
                break
            last_dash_row -= 1
            
        if last_dash_date:
            ws_ref = wb['GOLDBETA']
            new_dates = []
            for r in range(10, ws_ref.max_row + 1):
                d_val = ws_ref.cell(row=r, column=1).value
                if d_val:
                    d_parsed = None
                    if isinstance(d_val, datetime.datetime):
                        d_parsed = d_val.date()
                    else:
                        try:
                            d_parsed = pd.to_datetime(d_val).date()
                        except:
                            pass
                    if d_parsed and d_parsed > last_dash_date:
                        price = ws_ref.cell(row=r, column=3).value or 0
                        if float(price) > 0:
                            new_dates.append(d_parsed)
                            
            if new_dates:
                df_gold = get_historical_gold(new_dates[0], new_dates[-1])
                curr_row = ws_dash.max_row + 1
                for target_date in new_dates:
                    ws_dash.cell(row=curr_row, column=1, value=pd.to_datetime(target_date))
                    ws_dash.cell(row=curr_row, column=1).number_format = 'dd-mmm-yy'
                    from openpyxl.styles import Alignment
                    ws_dash.cell(row=curr_row, column=1).alignment = Alignment(horizontal="center", vertical="center")
                    
                    for etf_n, (c_close, c_val, units) in etfs_dash.items():
                        ws_etf = wb[etf_n]
                        etf_price = 0
                        for r in range(10, ws_etf.max_row + 1):
                            d_val = ws_etf.cell(row=r, column=1).value
                            if d_val:
                                d_parsed = None
                                if isinstance(d_val, datetime.datetime):
                                    d_parsed = d_val.date()
                                else:
                                    try:
                                        d_parsed = pd.to_datetime(d_val).date()
                                    except:
                                        pass
                                if d_parsed == target_date:
                                    etf_price = ws_etf.cell(row=r, column=3).value or 0
                                    break
                        
                        etf_value = float(etf_price) * units
                        cell_price = ws_dash.cell(row=curr_row, column=c_close)
                        cell_price.value = etf_price
                        cell_price.number_format = '[$₹-en-IN]#,##0.00'
                        
                        cell_val_obj = ws_dash.cell(row=curr_row, column=c_val)
                        cell_val_obj.value = etf_value
                        cell_val_obj.number_format = '[$₹-en-IN]#,##0.00'
                        
                    # Add Physical Gold Price
                    dt_pd = pd.to_datetime(target_date)
                    val_gold = 0
                    if dt_pd in df_gold.index:
                        val_gold = df_gold.loc[dt_pd, 'Domestic_10g']
                    else:
                        past_dates = df_gold[df_gold.index <= dt_pd]
                        if not past_dates.empty:
                            val_gold = past_dates.iloc[-1]['Domestic_10g']
                            
                    c_gold = ws_dash.cell(row=curr_row, column=12)
                    c_gold.value = float(val_gold)
                    c_gold.number_format = '[$₹-en-IN]#,##0.00'
                        
                    logging.info(f"[Dashboard] Appended new date {target_date} with Gold Price: {val_gold}")
                    curr_row += 1

    logging.info("Saving updated tracker...")
    try:
        wb.save(file_path)
        logging.info("Tracker saved successfully.")
    except Exception as e:
        logging.error(f"Error saving file: {e}")

if __name__ == "__main__":
    file_path = 'Gold_ETF_Daily_Price_Tracker.xlsx'
    update_tracker(file_path)
